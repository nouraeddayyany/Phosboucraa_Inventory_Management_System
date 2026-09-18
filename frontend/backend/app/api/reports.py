from datetime import datetime, timedelta
from typing import Any
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session
import io

from app.db.database import get_db
from app.models.user import User
from app.models.article import Article
from app.models.category import Category
from app.models.stock import Stock
from app.models.stock_movement import StockMovement, MovementType
from app.models.stock_request import StockRequest, RequestStatus
from app.utils.deps import check_permission
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

router = APIRouter()

@router.get("/dashboard")
def dashboard(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    total_articles = db.query(func.count(Article.id)).filter(Article.status == "ACTIVE").scalar() or 0
    total_stock = db.query(func.coalesce(func.sum(Stock.quantity), 0)).scalar() or 0
    critical = db.query(func.count(Stock.id)).filter(Stock.quantity <= 0).scalar() or 0
    low = (
        db.query(func.count(Stock.id))
        .join(Article, Article.id == Stock.article_id)
        .filter(Stock.quantity > 0, Stock.quantity < Article.stock_min)
        .scalar() or 0
    )
    pending = db.query(func.count(StockRequest.id)).filter(
        StockRequest.status.in_([RequestStatus.SUBMITTED, RequestStatus.PENDING_APPROVAL])
    ).scalar() or 0

    today = datetime.utcnow().date()
    start = datetime.combine(today, datetime.min.time())
    end = start + timedelta(days=1)
    counts = {}
    for movement_type in MovementType:
        counts[movement_type.value] = db.query(func.count(StockMovement.id)).filter(
            StockMovement.movement_type == movement_type,
            StockMovement.created_at >= start,
            StockMovement.created_at < end,
        ).scalar() or 0

    recent = db.query(StockMovement).order_by(StockMovement.created_at.desc()).limit(10).all()
    critical_rows = (
        db.query(Stock, Article)
        .join(Article, Article.id == Stock.article_id)
        .filter(Stock.quantity <= Article.stock_min)
        .order_by(Stock.quantity.asc()).limit(10).all()
    )
    return {
        "kpis": {
            "total_articles": total_articles,
            "total_stock": total_stock,
            "critical_stock": critical,
            "low_stock": low,
            "pending_requests": pending,
            "today_receipts": counts.get("RECEIPT", 0),
            "today_issues": counts.get("ISSUE", 0),
            "today_transfers": counts.get("TRANSFER", 0),
        },
        "movements_today": counts,
        "critical": [
            {"article_id": s.article_id, "quantity": s.quantity, "minimum_stock": a.stock_min}
            for s,a in critical_rows
        ],
        "recent_activity": [
            {"id": m.id, "movement_number": m.movement_number,
             "movement_type": m.movement_type, "article_id": m.article_id,
             "quantity": m.quantity, "user_id": m.user_id,
             "created_at": m.created_at}
            for m in recent
        ],
    }

@router.get("/stock")
def stock_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    rows = (
        db.query(Stock, Article, Category)
        .join(Article, Article.id == Stock.article_id)
        .outerjoin(Category, Category.id == Article.category_id)
        .all()
    )
    return [
        {"article_id": s.article_id, "code": a.code, "name": a.designation,
         "category": c.name if c else None, "location_id": s.location_id,
         "quantity": s.quantity, "minimum_stock": a.stock_min,
         "status": "CRITICAL" if s.quantity <= 0 else ("LOW" if s.quantity < a.stock_min else "NORMAL")}
        for s,a,c in rows
    ]


@router.get("/stock/export")
def stock_report_export(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    rows = (
        db.query(Stock, Article, Category)
        .join(Article, Article.id == Stock.article_id)
        .outerjoin(Category, Category.id == Article.category_id)
        .all()
    )

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Report"

    # Define headers
    headers = ["Article Code", "Article Name", "Category", "Location ID", "Quantity", "Minimum Stock", "Status"]
    ws.append(headers)

    # Style headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Add data rows
    for s, a, c in rows:
        status = "CRITICAL" if s.quantity <= 0 else ("LOW" if s.quantity < a.stock_min else "NORMAL")
        ws.append([
            a.code,
            a.designation,
            c.name if c else "N/A",
            str(s.location_id),
            s.quantity,
            a.stock_min,
            status
        ])

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"stock_report_{timestamp}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/movements")
def movements_report(
    date_from: datetime | None = Query(None),
    date_to: datetime | None = Query(None),
    movement_type: MovementType | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    q = db.query(StockMovement)
    if date_from: q = q.filter(StockMovement.created_at >= date_from)
    if date_to: q = q.filter(StockMovement.created_at <= date_to)
    if movement_type: q = q.filter(StockMovement.movement_type == movement_type)
    return q.order_by(StockMovement.created_at.desc()).limit(500).all()


@router.get("/movements/export")
def movements_report_export(
    date_from: datetime | None = Query(None),
    date_to: datetime | None = Query(None),
    movement_type: MovementType | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    q = db.query(StockMovement)
    if date_from: q = q.filter(StockMovement.created_at >= date_from)
    if date_to: q = q.filter(StockMovement.created_at <= date_to)
    if movement_type: q = q.filter(StockMovement.movement_type == movement_type)
    movements = q.order_by(StockMovement.created_at.desc()).limit(500).all()

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movements Report"

    # Define headers
    headers = ["Movement Number", "Movement Type", "Article ID", "Quantity", "Location ID", "Source Location", "Destination Location", "Reason", "Reference", "User ID", "Created At"]
    ws.append(headers)

    # Style headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Add data rows
    for m in movements:
        ws.append([
            m.movement_number,
            m.movement_type.value,
            str(m.article_id),
            m.quantity,
            str(m.location_id) if m.location_id else "N/A",
            str(m.source_location_id) if m.source_location_id else "N/A",
            str(m.destination_location_id) if m.destination_location_id else "N/A",
            m.reason or "N/A",
            m.reference or "N/A",
            str(m.user_id),
            m.created_at.strftime("%Y-%m-%d %H:%M:%S")
        ])

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"movements_report_{timestamp}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/critical-stock")
def critical_stock(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    return dashboard(db, current_user)["critical"]

@router.get("/inventory")
def inventory_report(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    from app.models.inventory import Inventory
    return db.query(Inventory).order_by(Inventory.created_at.desc()).all()


@router.get("/stock/export/pdf")
def stock_report_export_pdf(
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    rows = (
        db.query(Stock, Article, Category)
        .join(Article, Article.id == Stock.article_id)
        .outerjoin(Category, Category.id == Article.category_id)
        .all()
    )

    # Create PDF buffer
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#4472C4'),
        spaceAfter=30,
    )
    elements.append(Paragraph("Stock Report", title_style))
    elements.append(Spacer(1, 12))

    # Date
    date_style = ParagraphStyle(
        'CustomDate',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.grey,
    )
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", date_style))
    elements.append(Spacer(1, 20))

    # Table data
    table_data = [['Article Code', 'Article Name', 'Category', 'Location', 'Quantity', 'Min Stock', 'Status']]
    for s, a, c in rows:
        status = "CRITICAL" if s.quantity <= 0 else ("LOW" if s.quantity < a.stock_min else "NORMAL")
        table_data.append([
            a.code,
            a.designation,
            c.name if c else "N/A",
            str(s.location_id),
            str(s.quantity),
            str(a.stock_min),
            status,
        ])

    # Create table
    table = Table(table_data, colWidths=[1.5*inch, 2.5*inch, 1.5*inch, 1*inch, 0.8*inch, 0.8*inch, 1*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))
    elements.append(table)

    # Build PDF
    doc.build(elements)
    output.seek(0)

    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"stock_report_{timestamp}.pdf"

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/movements/export/pdf")
def movements_report_export_pdf(
    date_from: datetime | None = Query(None),
    date_to: datetime | None = Query(None),
    movement_type: MovementType | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(check_permission("REPORT_READ")),
):
    q = db.query(StockMovement)
    if date_from: q = q.filter(StockMovement.created_at >= date_from)
    if date_to: q = q.filter(StockMovement.created_at <= date_to)
    if movement_type: q = q.filter(StockMovement.movement_type == movement_type)
    movements = q.order_by(StockMovement.created_at.desc()).limit(500).all()

    # Create PDF buffer
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    elements = []
    styles = getSampleStyleSheet()

    # Title
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#4472C4'),
        spaceAfter=30,
    )
    elements.append(Paragraph("Movements Report", title_style))
    elements.append(Spacer(1, 12))

    # Date
    date_style = ParagraphStyle(
        'CustomDate',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.grey,
    )
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", date_style))
    elements.append(Spacer(1, 20))

    # Table data
    table_data = [['Movement #', 'Type', 'Article ID', 'Quantity', 'Location', 'Source', 'Destination', 'Reason', 'User', 'Date']]
    for m in movements:
        table_data.append([
            m.movement_number,
            m.movement_type.value,
            str(m.article_id),
            str(m.quantity),
            str(m.location_id) if m.location_id else "N/A",
            str(m.source_location_id) if m.source_location_id else "N/A",
            str(m.destination_location_id) if m.destination_location_id else "N/A",
            m.reason or "N/A",
            str(m.user_id),
            m.created_at.strftime("%Y-%m-%d %H:%M:%S"),
        ])

    # Create table
    table = Table(table_data, colWidths=[1.2*inch, 1*inch, 1*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch, 1.2*inch, 0.8*inch, 1.2*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    elements.append(table)

    # Build PDF
    doc.build(elements)
    output.seek(0)

    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"movements_report_{timestamp}.pdf"

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
